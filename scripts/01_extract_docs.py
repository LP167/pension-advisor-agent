# -*- coding: utf-8 -*-
"""docs_renamed 58종 전량 텍스트 추출 + 유형 판별.
출력: data/raw_text/<docid>.json  {doc_id, ext, n_pages, is_image_pdf, pages:[{page,text}], tables:[...]}
"""
import os, sys, json, re, traceback
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC  = os.path.join(os.path.dirname(BASE), "2.연금", "docs_renamed")
OUT  = os.path.join(BASE, "data", "raw_text")
os.makedirs(OUT, exist_ok=True)

def norm(s):
    if not s: return ""
    s = s.replace(" "," ")
    s = re.sub(r"[ \t]+"," ", s)
    s = re.sub(r"\n{3,}","\n\n", s)
    return s.strip()

def do_pdf(p):
    import pdfplumber
    pages, tables = [], []
    with pdfplumber.open(p) as pdf:
        for i, pg in enumerate(pdf.pages, 1):
            t = norm(pg.extract_text() or "")
            pages.append({"page": i, "text": t})
            try:
                for ti, tb in enumerate(pg.extract_tables() or []):
                    tables.append({"page": i, "idx": ti, "rows": tb})
            except Exception: pass
    chars = sum(len(x["text"]) for x in pages)
    return pages, tables, chars

def do_docx(p):
    import docx
    d = docx.Document(p)
    paras = [norm(x.text) for x in d.paragraphs]
    paras = [x for x in paras if x]
    tables = []
    for ti, tb in enumerate(d.tables):
        rows = [[norm(c.text) for c in r.cells] for r in tb.rows]
        tables.append({"page": None, "idx": ti, "rows": rows})
    return [{"page": 1, "text": "\n".join(paras)}], tables, sum(len(x) for x in paras)

def do_xlsx(p):
    import openpyxl
    wb = openpyxl.load_workbook(p, data_only=True)
    tables, txt = [], []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(["" if c is None else str(c) for c in r])
        tables.append({"sheet": ws.title, "idx": 0, "rows": rows})
        txt.append(ws.title)
    return [{"page": 1, "text": "\n".join(txt)}], tables, 1
def do_pptx(p):
    from pptx import Presentation
    pr = Presentation(p); pages=[]
    for i, sl in enumerate(pr.slides, 1):
        buf=[]
        for sh in sl.shapes:
            if sh.has_text_frame: buf.append(norm(sh.text_frame.text))
        pages.append({"page": i, "text": "\n".join([b for b in buf if b])})
    return pages, [], sum(len(x["text"]) for x in pages)

res=[]
files = sorted(os.listdir(SRC), key=lambda f:(f.split('.')[-1], int(re.sub(r'\D','',f) or 0)))
for f in files:
    if f.startswith('.'): continue
    p = os.path.join(SRC, f); doc_id, ext = os.path.splitext(f); ext = ext.lower().lstrip('.')
    try:
        if ext=="pdf":   pages, tables, chars = do_pdf(p)
        elif ext=="docx":pages, tables, chars = do_docx(p)
        elif ext=="xlsx":pages, tables, chars = do_xlsx(p)
        elif ext=="pptx":pages, tables, chars = do_pptx(p)
        else: continue
        n=len(pages); per = chars/max(n,1)
        is_img = (ext=="pdf" and per < 80)
        json.dump({"doc_id":doc_id,"ext":ext,"n_pages":n,"chars":chars,
                   "is_image_pdf":is_img,"pages":pages,"tables":tables},
                  open(os.path.join(OUT,doc_id+".json"),"w"), ensure_ascii=False)
        res.append((doc_id, ext, n, chars, is_img))
        print(f"{doc_id}\t{ext}\tp={n}\tchars={chars}\timg={is_img}", flush=True)
    except Exception as e:
        print(f"{doc_id}\tERROR\t{e}", flush=True); traceback.print_exc()
img=[r[0] for r in res if r[4]]
json.dump({"total":len(res),"image_pdfs":img}, open(os.path.join(BASE,"data","_extract_summary.json"),"w"), ensure_ascii=False, indent=1)
print("\nTOTAL", len(res), "IMAGE_PDFS", len(img), img)
